import openpyxl
from fastapi import UploadFile
from langchain_core.documents import Document

class XLSXLoader:
    def __init__(self, files: list[UploadFile]):
        self.files = files

    def load(self) -> list:
        documents = []

        for upload_file in self.files:
            full_text = []
            wb = openpyxl.load_workbook(upload_file.file)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    full_text.append(", ".join([str(cell) for cell in row]))
            content = "\n".join(full_text)
            metadata = {"source": upload_file.filename}
            doc = Document(page_content=content, metadata=metadata)
            documents.append(doc)

        return documents
